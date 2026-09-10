"""
Genera UN archivo .xlsx POR SECCIÓN (course_id) — ya no se mezclan los
alumnos de D16 y D23 en el mismo Excel.
"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from . import config, db

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FLAG_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
FAIL_FILL = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")


def _style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autofit(ws):
    for col_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 10), 45)


def _safe_filename(name: str) -> str:
    return re.sub(r'[\\/*?:"<>|]', "_", name).strip()


def build_workbook_for_course(cfg: dict, course_id: str) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    with db.get_conn() as conn:
        students = conn.execute(
            "SELECT student_id, full_name FROM students WHERE course_id=? ORDER BY full_name",
            (course_id,),
        ).fetchall()

        # --- Resumen ---
        ws = wb.create_sheet("Resumen")
        categories = list(cfg["weights"].keys())
        ws.append(["Alumno", *[f"{c} ({int(cfg['weights'][c]*100)}%)" for c in categories], "Final"])
        _style_header(ws)

        finals = db.compute_final_grades(conn, cfg["weights"], course_id=course_id)
        for s in students:
            data = finals.get(s["student_id"], {"final": 0.0, "breakdown": {}})
            row = [s["full_name"]]
            for c in categories:
                pct = data["breakdown"].get(c)
                row.append(round(pct * 100, 1) if pct is not None else "—")
            row.append(data["final"])
            ws.append(row)
        _autofit(ws)

        # --- Una hoja por categoría, formato pivote: alumno x tarea (solo de esta sección) ---
        for category in categories:
            ws = wb.create_sheet(category[:31])

            courseworks = conn.execute(
                "SELECT coursework_id, title, unidad FROM coursework "
                "WHERE category = ? AND course_id = ? ORDER BY unidad, title",
                (category, course_id),
            ).fetchall()

            headers = ["Alumno"] + [cw["title"] for cw in courseworks] + ["Promedio"]
            ws.append(headers)
            _style_header(ws)

            scores = {}
            rows = conn.execute(
                """
                SELECT s.student_id, s.coursework_id, g.score_raw, cw.max_points, g.plagiarism_score
                FROM grades g
                JOIN submissions s ON s.submission_id = g.submission_id
                JOIN coursework cw ON cw.coursework_id = s.coursework_id
                WHERE cw.category = ? AND cw.course_id = ?
                """,
                (category, course_id),
            ).fetchall()
            for r in rows:
                scores[(r["student_id"], r["coursework_id"])] = (r["score_raw"], r["max_points"], r["plagiarism_score"])

            for st in students:
                row_idx = ws.max_row + 1
                row_values = [st["full_name"]]
                pcts = []
                for cw in courseworks:
                    entry = scores.get((st["student_id"], cw["coursework_id"]))
                    if entry is None:
                        row_values.append("—")
                    else:
                        score_raw, max_points, plag = entry
                        row_values.append(score_raw)
                        pcts.append(score_raw / max_points if max_points else 0)
                promedio = round(sum(pcts) / len(pcts) * 100, 1) if pcts else "—"
                row_values.append(promedio)
                ws.append(row_values)

                for col_idx, cw in enumerate(courseworks, start=2):
                    entry = scores.get((st["student_id"], cw["coursework_id"]))
                    if entry and entry[2] and entry[2] >= cfg["plagiarism"]["similarity_threshold"]:
                        ws.cell(row=row_idx, column=col_idx).fill = FLAG_FILL
            _autofit(ws)

        # --- Detalle por problema (solo de esta sección) ---
        ws = wb.create_sheet("Detalle_Problemas")
        ws.append(["Alumno", "Tarea", "Problema", "Compiló", "Tests", "Puntos", "Feedback"])
        _style_header(ws)

        rows = conn.execute(
            """
            SELECT st.full_name, cw.title, pg.problem_index, pg.compile_ok,
                   pg.tests_passed, pg.tests_total, pg.score_raw, pg.max_points, pg.feedback
            FROM problem_grades pg
            JOIN submissions s ON s.submission_id = pg.submission_id
            JOIN students st ON st.student_id = s.student_id
            JOIN coursework cw ON cw.coursework_id = s.coursework_id
            WHERE st.course_id = ?
            ORDER BY cw.unidad, st.full_name, pg.problem_index
            """,
            (course_id,),
        ).fetchall()

        for r in rows:
            row_idx = ws.max_row + 1
            ws.append([
                r["full_name"], r["title"], f"p{r['problem_index']}",
                "Sí" if r["compile_ok"] else "No",
                f"{r['tests_passed']}/{r['tests_total']}" if r["tests_total"] else "—",
                f"{r['score_raw']}/{r['max_points']}",
                (r["feedback"] or "")[:200],
            ])
            if r["score_raw"] == 0:
                for cell in ws[row_idx]:
                    cell.fill = FAIL_FILL
        _autofit(ws)

    return wb


def export():
    """Genera un .xlsx por cada sección que tenga alumnos registrados.
    Regresa la lista de rutas generadas."""
    cfg = config.load_config()
    config.ensure_dirs()

    with db.get_conn() as conn:
        course_ids = db.get_all_course_ids(conn)
        course_names = {cid: db.get_course_name(conn, cid) for cid in course_ids}

    paths = []
    for course_id in course_ids:
        wb = build_workbook_for_course(cfg, course_id)
        nombre_archivo = f"calificaciones_{_safe_filename(course_names[course_id])}_{course_id[-6:]}.xlsx"
        path = config.DATA_DIR / nombre_archivo
        wb.save(path)
        paths.append(path)
        print(f"  Generado: {path}")

    return paths
